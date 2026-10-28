"""Builds the Minutes of the Meeting Word file in the Telex template layout (see MM.docx)."""
import io
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Image
from docx.oxml import OxmlElement
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

# ---- letterhead: edit these to change the company block on every page
COMPANY_NAME = "TELEX BUSINESS SUPPORT SERVICES INC."
COMPANY_LINES = [
    "Telex Building, National Highway",
    "Brgy. Cawayan Bugtong, Guimba, Nueva Ecija 3115",
    "044-950-4196",
]
LOGO = Path(__file__).parent / "logo.png"

BODY_FONT = "Leelawadee"   # Word falls back automatically if the font is not installed
HEAD_FONT = "Leelawadee"
BODY_PT = 12
J = WD_ALIGN_PARAGRAPH.JUSTIFY


def _font(run, name=BODY_FONT, size=BODY_PT, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def _para(doc, text="", bold=False, size=BODY_PT, align=None, after=8, before=0, font=BODY_FONT, keep_next=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after, pf.space_before, pf.line_spacing = Pt(after), Pt(before), 1.15
    pf.keep_with_next = keep_next
    if align is not None:
        p.alignment = align
    if text:
        _font(p.add_run(text), font, size, bold)
    return p


def _label_line(doc, label, value):
    p = _para(doc, after=0)
    _font(p.add_run(label), bold=True)
    _font(p.add_run(f": {value}"))
    return p


def _rule(doc, thick=False):
    """Horizontal line between sections. Runs a little past the text margins, like the template."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(6), Pt(8)
    pf.left_indent = pf.right_indent = Inches(-0.3)
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    sz, color = ("24", "BFBFBF") if thick else ("6", "7F7F7F")
    for k, v in (("val", "single"), ("sz", sz), ("space", "1"), ("color", color)):
        bottom.set(qn(f"w:{k}"), v)
    bdr.append(bottom)
    pPr.append(bdr)
    return p


def _heading(doc, text, after=4):
    return _para(doc, text, bold=True, after=after, before=2, keep_next=True, align=J)


def _bullet(doc, text, level):
    sym, left = ("▪", 0.5) if level == 1 else ("•", 1.0)
    p = _para(doc, after=1)
    p.paragraph_format.left_indent = Inches(left)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    _font(p.add_run(f"{sym}  "), "Arial", BODY_PT - 1)
    _font(p.add_run(str(text)))
    return p


def _no_borders(table):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    table._tbl.tblPr.append(borders)


def _header(section):
    section.header_distance = Inches(0.6)
    hdr = section.header
    tbl = hdr.add_table(rows=1, cols=2, width=Inches(6.5))
    _no_borders(tbl)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    left, right = tbl.rows[0].cells
    left.width, right.width = Inches(4.8), Inches(1.7)

    p = left.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    _font(p.add_run(COMPANY_NAME), HEAD_FONT, 11, True)
    for line in COMPANY_LINES:
        q = left.add_paragraph()
        q.paragraph_format.space_after = Pt(0)
        _font(q.add_run(line), HEAD_FONT, 11)

    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if LOGO.exists():
        rp.add_run().add_picture(str(LOGO), width=Inches(1.2))

    # python-docx puts an empty paragraph first in a header; drop it and keep a tiny one after the table
    first = hdr.paragraphs[0]
    first._p.getparent().remove(first._p)
    tail = hdr.add_paragraph()
    tail.paragraph_format.space_after = Pt(0)
    _font(tail.add_run(""), size=2)


def _shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _title_and_info(d, m):
    _para(d, "Minutes of the Meeting", bold=True, size=14, after=0)
    _para(d, m.get("title") or "Meeting", bold=True, size=14, after=12)
    _label_line(d, "Date", m.get("date") or "TBD")
    _label_line(d, "Time", m.get("time") or "TBD")
    _label_line(d, "Location", m.get("location") or "TBD")
    _label_line(d, "Purpose", m.get("purpose") or "TBD")
    _para(d, "Attendees:", bold=True, after=2)
    for g in m.get("attendees") or []:
        name = (g.get("group") or "").strip()
        members = g.get("members") or []
        if name:
            _bullet(d, name, 1)
            for x in members:
                _bullet(d, x, 2)
        else:
            for x in members:
                _bullet(d, x, 1)
    _rule(d)


def _trim_signature(blob):
    """Crop the empty margin around a signature and make a white/transparent background transparent.
    A photo or scan has a lot of blank paper around the ink, which would make the signature look tiny."""
    try:
        from PIL import Image as PILImage
        im = PILImage.open(io.BytesIO(blob)).convert("RGBA")
        px = im.load()
        w, h = im.size
        xs, ys = [], []
        for y in range(h):
            for x in range(w):
                r, g, b, a = px[x, y]
                if a > 40 and (r < 215 or g < 215 or b < 215):   # ink: opaque and not near-white
                    xs.append(x)
                    ys.append(y)
                else:
                    px[x, y] = (255, 255, 255, 0)                 # background -> transparent
        if not xs:
            return blob
        pad = 4
        box = (max(min(xs) - pad, 0), max(min(ys) - pad, 0), min(max(xs) + pad + 1, w), min(max(ys) + pad + 1, h))
        out = io.BytesIO()
        im.crop(box).save(out, format="PNG")
        return out.getvalue()
    except Exception:
        return blob   # Pillow missing or unreadable image: use it as it is


SIG_MAX_W, SIG_MAX_H = 1.2, 0.46   # inches, applied after the empty margins around the signature are trimmed


def signature_size(blob):
    """(width, height) in inches for a trimmed signature image. Also used by the on-screen preview."""
    img = Image.from_blob(blob)
    scale = min(SIG_MAX_W / (img.px_width / 96), SIG_MAX_H / (img.px_height / 96), 12)
    return img.px_width / 96 * scale, img.px_height / 96 * scale


def _float_picture(par, blob, width_in, x_in, y_in):
    """Insert a picture that floats behind the text (so a white photo background never hides the name), positioned relative to the paragraph."""
    run = par.add_run()
    inline = run.add_picture(io.BytesIO(blob), width=Inches(width_in))._inline
    cx, cy = inline.extent.cx, inline.extent.cy
    anchor = parse_xml(
        f'<wp:anchor {nsdecls("wp", "a", "pic", "r")} distT="0" distB="0" distL="0" distR="0" simplePos="0" '
        f'relativeHeight="251658240" behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">'
        f'<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="column"><wp:posOffset>{int(Inches(x_in))}</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="paragraph"><wp:posOffset>{int(Inches(y_in))}</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{cx}" cy="{cy}"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
        f'<wp:docPr id="9001" name="E-signature"/><wp:cNvGraphicFramePr/></wp:anchor>')
    anchor.append(inline.graphic)
    drawing = inline.getparent()
    drawing.remove(inline)
    drawing.append(anchor)


def _signature(d, m):
    """'Prepared By' block. If an e-signature image was provided it sits right above the printed name."""
    pb = m.get("prepared_by") or {}
    _para(d, "Prepared By:", after=0, before=14, keep_next=True)
    sig = _trim_signature(m.get("signature")) if m.get("signature") else None
    name = _para(d, pb.get("name") or "______________________", after=0, before=14, keep_next=True)
    if sig:
        try:
            w_in, h_in = signature_size(sig)
            overlap = min(0.24, h_in * 0.25)   # how far the signature dips onto the printed name
            name.paragraph_format.space_before = Pt(max(14, (h_in - overlap) * 72))
            _float_picture(name, sig, w_in, x_in=0.1, y_in=0)   # offset is measured from the top of the spacing above the name
        except Exception:
            name.paragraph_format.space_before = Pt(30)   # unreadable image: leave blank signing space
    role = _para(d, pb.get("role") or "", after=0)
    role.paragraph_format.left_indent = Inches(0.1)
    return name


def _hearing_sections(d, m):
    """Admin hearing: 1 Call to Order, 2 Purpose, 3 Discussion Summary, 4 Management Review and Decision, 5 Handbook."""
    _heading(d, "1. Call to Order")
    _para(d, m.get("call_to_order") or "The meeting was called to order.", align=J)
    _rule(d)

    _heading(d, "2. Purpose of the Meeting")
    _para(d, m.get("summary") or "TBD", align=J)
    _rule(d)

    _heading(d, "3. Discussion Summary")
    for para in m.get("discussion") or ["TBD"]:
        _para(d, str(para), align=J)
    _rule(d, thick=True)

    _heading(d, "4. Management Review and Decision")
    if (m.get("review_intro") or "").strip():
        _para(d, m["review_intro"].strip(), align=J)
    _para(d, "Decision:", bold=True, after=6)
    for x in m.get("decisions") or ["TBD"]:
        _para(d, str(x), bold=True, after=6)

    offense = (m.get("offense") or "").strip()
    handbook = [x.strip() for x in (m.get("handbook") or "").splitlines() if x.strip()]
    if offense or handbook:
        _rule(d, thick=True)
        if offense:
            _heading(d, f"5. {offense}", after=6)
        else:
            _heading(d, "5. Company Employee Handbook", after=6)
        _para(d, "Based on the Company Employee Handbook:", align=J)
        for line in handbook:
            _para(d, line, align=J)


def _general_sections(d, m):
    _heading(d, "1. Call to Order")
    _para(d, m.get("call_to_order") or "The meeting was called to order.", align=J)
    _rule(d)
    _heading(d, "2. Purpose of the Meeting")
    _para(d, m.get("summary") or "TBD", align=J)
    _rule(d)
    _heading(d, "3. Discussion Summary")
    for para in m.get("discussion") or ["TBD"]:
        _para(d, str(para), align=J)
    _rule(d, thick=True)
    _heading(d, "4. Decisions")
    _para(d, "Decision:", bold=True, after=4)
    for x in m.get("decisions") or ["No decision recorded."]:
        _para(d, str(x), bold=True, after=4)

    last = 4
    actions = m.get("action_items") or []
    if actions:
        last += 1
        _rule(d, thick=True)
        _heading(d, f"{last}. Action Items", after=6)
        t = d.add_table(rows=1, cols=3)
        t.style = "Table Grid"
        for i, (h, w) in enumerate((("Task", 3.5), ("Owner", 1.5), ("Deadline", 1.5))):
            c = t.rows[0].cells[i]
            c.width = Inches(w)
            _shade(c, "EDEDED")
            _font(c.paragraphs[0].add_run(h), size=BODY_PT - 1, bold=True)
        for a in actions:
            cells = t.add_row().cells
            for i, (key, dflt, w) in enumerate((("task", "", 3.5), ("owner", "TBD", 1.5), ("deadline", "TBD", 1.5))):
                cells[i].width = Inches(w)
                _font(cells[i].paragraphs[0].add_run(str(a.get(key) or dflt)), size=BODY_PT - 1)
        _para(d, after=2)

    notes = (m.get("notes") or "").strip()
    if notes:
        last += 1
        _rule(d, thick=True)
        _heading(d, f"{last}. Additional Notes", after=6)
        for para in notes.split("\n"):
            if para.strip():
                _para(d, para.strip(), align=J)


def build_docx(m, include_transcript=True):
    """m['kind'] is 'hearing' (HR admin hearing, as in MM.docx) or 'general'."""
    d = Document()
    normal = d.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(BODY_PT)

    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(14)   # Legal, same as the template
    sec.left_margin = sec.right_margin = Inches(1)
    sec.top_margin, sec.bottom_margin = Inches(2.1), Inches(0.9)
    _header(sec)

    _title_and_info(d, m)
    if m.get("kind") == "hearing":
        _hearing_sections(d, m)
    else:
        _general_sections(d, m)
    _signature(d, m)

    if include_transcript and m.get("transcript"):
        d.add_page_break()
        _para(d, "Appendix: Conversation Transcript", bold=True, size=14, after=8)
        for line in str(m["transcript"]).split("\n"):
            if not line.strip():
                continue
            mm = re.match(r"^([^:\n]{1,30}):\s*(.*)$", line.strip())
            p = _para(d, after=4)
            if mm:
                _font(p.add_run(mm.group(1) + ": "), size=BODY_PT - 1, bold=True)
                _font(p.add_run(mm.group(2)), size=BODY_PT - 1)
            else:
                _font(p.add_run(line.strip()), size=BODY_PT - 1)

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
